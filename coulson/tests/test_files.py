import pytest

from coulson.config import load
from coulson.memory import Memory
from coulson.tools import Context, files, load_all
from coulson.tools.system import _keys_risk, parse_keys


@pytest.fixture
def env(tmp_path, monkeypatch):
    home = tmp_path / "home"
    (home / "Desktop").mkdir(parents=True)
    monkeypatch.setattr(files, "HOME", home)
    monkeypatch.setenv("HOME", str(home))
    cfg = load()
    cfg["files"] = {"workspace": str(home / "Documents" / "Колсон")}
    asked = []
    ctx = Context(cfg, memory=Memory(tmp_path / "m.db"), confirm=lambda d: asked.append(d) or False)
    return load_all(), ctx, home, asked


def test_resolve(env):
    _, ctx, home, _ = env
    assert files.resolve(ctx, "page.html") == (home / "Documents" / "Колсон" / "page.html").resolve()
    assert files.resolve(ctx, "Рабочий стол/отчёт.html") == (home / "Desktop" / "отчёт.html").resolve()
    assert files.resolve(ctx, "~/Desktop/a.txt") == (home / "Desktop" / "a.txt").resolve()


def test_write_file_new_is_auto_overwrite_asks(env):
    reg, ctx, home, asked = env
    res = reg.execute("write_file", {"path": "hello.html", "content": "<h1>Привет</h1>", "open": False}, ctx)
    target = home / "Documents" / "Колсон" / "hello.html"
    assert "Создан файл" in res and target.read_text(encoding="utf-8") == "<h1>Привет</h1>" and not asked
    res = reg.execute("write_file", {"path": "hello.html", "content": "new", "open": False}, ctx)
    assert "НЕ подтвердил" in res and asked and target.read_text(encoding="utf-8") == "<h1>Привет</h1>"


def test_write_file_hidden_or_outside_asks(env):
    reg, ctx, home, asked = env
    assert "НЕ подтвердил" in reg.execute("write_file", {"path": "~/.zshrc", "content": "x", "open": False}, ctx)
    assert "НЕ подтвердил" in reg.execute("write_file", {"path": "/etc/hosts", "content": "x", "open": False}, ctx)
    assert len(asked) == 2


def test_create_table_xlsx_with_chart(env):
    from openpyxl import load_workbook
    reg, ctx, home, asked = env
    res = reg.execute("create_table", {
        "title": "Загрузка диска", "columns": ["Папка", "Размер, ГБ"],
        "rows": [["Library", 48.2], ["Downloads", "12,5"], ["Movies", 7]], "chart": "pie", "open": False}, ctx)
    target = home / "Documents" / "Колсон" / "Загрузка диска.xlsx"
    assert "Таблица сохранена" in res and target.exists()
    ws = load_workbook(target).active
    assert ws["A1"].value == "Папка" and ws["B3"].value == 12.5 and len(ws._charts) == 1
    assert "Library | 48.2" in reg.execute("read_file", {"path": "Загрузка диска.xlsx"}, ctx)


def test_create_table_csv(env):
    reg, ctx, home, _ = env
    reg.execute("create_table", {"title": "t", "columns": ["a", "b"], "rows": [[1, 2]], "path": "Рабочий стол/t.csv",
                                 "open": False}, ctx)
    assert (home / "Desktop" / "t.csv").read_text(encoding="utf-8-sig").splitlines() == ["a;b", "1;2"]


def test_disk_usage_and_list_dir(env):
    reg, ctx, home, _ = env
    (home / "Movies").mkdir()
    (home / "Movies" / "big.bin").write_bytes(b"0" * 3_000_000)
    out = reg.execute("disk_usage", {"path": "~"}, ctx)
    assert out.startswith("Диск: всего") and "Movies" in out.splitlines()[2]
    assert "big.bin" in reg.execute("list_dir", {"path": "~/Movies"}, ctx)


def test_keys():
    assert parse_keys("cmd+shift+t") == ('keystroke "t"', ["command down", "shift down"])
    assert parse_keys("esc") == ("key code 53", [])
    with pytest.raises(ValueError):
        parse_keys("cmd+банан")
    assert _keys_risk(None, "cmd+q") and _keys_risk(None, "cmd+delete")
    assert _keys_risk(None, "cmd+t") is None and _keys_risk(None, "space") is None
