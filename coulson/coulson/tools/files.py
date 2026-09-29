"""Файлы: создать (текст, HTML, код), таблица Excel с диаграммой, прочитать, папки, загрузка диска."""
from __future__ import annotations

import csv
import os
import re
import subprocess
import time
from pathlib import Path

from . import registry, run

HOME = Path.home()

# Как пользователь (или модель) может назвать стандартные папки
_FOLDERS = {
    "рабочий стол": "~/Desktop", "рабочем столе": "~/Desktop", "рабочего стола": "~/Desktop", "desktop": "~/Desktop",
    "документы": "~/Documents", "документах": "~/Documents", "documents": "~/Documents",
    "загрузки": "~/Downloads", "загрузках": "~/Downloads", "downloads": "~/Downloads",
    "изображения": "~/Pictures", "картинки": "~/Pictures", "pictures": "~/Pictures",
    "музыка": "~/Music", "music": "~/Music", "видео": "~/Movies", "фильмы": "~/Movies", "movies": "~/Movies",
    "домашняя папка": "~", "home": "~",
}


def workspace(ctx) -> Path:
    ws = Path(os.path.expanduser(ctx.cfg.get("files", {}).get("workspace", "~/Documents/Колсон")))
    ws.mkdir(parents=True, exist_ok=True)
    return ws


def resolve(ctx, path: str) -> Path:
    """Относительный путь — в папку Колсона; «Рабочий стол/файл.html» — на рабочий стол."""
    p = (path or "").strip().strip("\"'«»")
    low = p.lower()
    for alias in sorted(_FOLDERS, key=len, reverse=True):
        if low == alias or low.startswith(alias + "/"):
            p = _FOLDERS[alias] + p[len(alias):]
            break
    p = os.path.expanduser(p)
    if not os.path.isabs(p):
        p = str(workspace(ctx) / p)
    return Path(p).resolve()


def _write_risk(target: Path) -> str | None:
    home = HOME.resolve()
    if home not in target.parents:
        return f"запись вне домашней папки: {target}"
    rel = target.relative_to(home)
    if rel.parts and (rel.parts[0] in ("Library", "Applications") or any(part.startswith(".") for part in rel.parts)):
        return f"изменение системного или скрытого файла {target}"
    if target.exists():
        return f"перезаписать существующий файл {target.name}"
    return None


def _open(ctx, target: Path) -> str:
    if target.suffix.lower() in (".html", ".htm", ".svg"):
        res = run(["open", "-a", ctx.cfg.browser.app, str(target)])
        if res == "OK":
            return res
    return run(["open", str(target)])


def _human(kb: float) -> str:
    for unit, div in (("ТБ", 1024 ** 3), ("ГБ", 1024 ** 2), ("МБ", 1024)):
        if kb >= div:
            return f"{kb / div:.1f} {unit}"
    return f"{kb:.0f} КБ"


# ---------------------------------------------------------------- создание

@registry.add("write_file", "Create a text file (HTML page, md/txt doc, code, JSON, SVG) with COMPLETE content. "
              "Relative path = ~/Documents/Колсон; 'Рабочий стол/x.html' = Desktop. Opens it after.",
              {"path": ("string", "name with extension"), "content": ("string", ""), "open": ("boolean", "")},
              ["path", "content"],
              risk=lambda ctx, path, content="", open=True: _write_risk(resolve(ctx, path)))
def write_file(ctx, path: str, content: str, open: bool = True) -> str:
    target = resolve(ctx, path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(content, encoding="utf-8")
    opened = f", открыл" if open and _open(ctx, target) == "OK" else ""
    return f"Создан файл {target} ({len(content)} символов){opened}"


def _cell(value):
    """'12,5' -> 12.5, чтобы в таблице были числа, а не текст."""
    if isinstance(value, str):
        v = value.strip().replace(" ", "").replace(" ", "")
        if re.fullmatch(r"[-+]?\d+([.,]\d+)?", v):
            return float(v.replace(",", ".")) if re.search(r"[.,]", v) else int(v)
    return value


@registry.add("create_table", "Create an Excel (.xlsx) or .csv table, optionally with a chart "
              "(labels = 1st column, values = 2nd).",
              {"title": ("string", ""), "columns": {"type": "array", "items": {"type": "string"}},
               "rows": {"type": "array", "items": {"type": "array", "items": {}}, "description": "numbers as numbers"},
               "path": ("string", "optional; .xlsx or .csv"), "chart": ("string", ""), "open": ("boolean", "")},
              ["title", "columns", "rows"], enums={"chart": ["none", "bar", "pie", "line"]},
              risk=lambda ctx, title, columns=None, rows=None, path="", chart="none", open=True:
                  _write_risk(_table_path(ctx, title, path)))
def create_table(ctx, title: str, columns: list, rows: list, path: str = "", chart: str = "none",
                 open: bool = True) -> str:
    target = _table_path(ctx, title, path)
    target.parent.mkdir(parents=True, exist_ok=True)
    rows = [[_cell(v) for v in (r if isinstance(r, list) else [r])] for r in rows]

    if target.suffix.lower() == ".csv":
        with target.open("w", newline="", encoding="utf-8-sig") as f:  # BOM — чтобы Excel понял кириллицу
            w = csv.writer(f, delimiter=";")
            w.writerow(columns)
            w.writerows(rows)
    else:
        _write_xlsx(target, title, columns, rows, chart)
    opened = ", открыл" if open and _open(ctx, target) == "OK" else ""
    return f"Таблица сохранена: {target} ({len(rows)} строк){opened}"


def _table_path(ctx, title: str, path: str) -> Path:
    if path:
        target = resolve(ctx, path)
        return target if target.suffix.lower() in (".xlsx", ".csv") else target.with_suffix(".xlsx")
    name = re.sub(r"[^\w\s-]", "", title).strip()[:60] or f"Таблица {time.strftime('%Y-%m-%d %H-%M')}"
    return resolve(ctx, f"{name}.xlsx")


def _write_xlsx(target: Path, title: str, columns: list, rows: list, chart: str) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Лист1"
    ws.append(list(columns))
    for r in rows:
        ws.append(r)
    head = PatternFill("solid", fgColor="1E6BFF")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, 1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows if len(r) >= i]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"

    if chart in ("bar", "pie", "line") and len(columns) >= 2 and rows:
        ch = {"bar": BarChart, "pie": PieChart, "line": LineChart}[chart]()
        ch.title = title
        data = Reference(ws, min_col=2, min_row=1, max_row=len(rows) + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=len(rows) + 1)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ch.height, ch.width = 9, 16
        ws.add_chart(ch, f"{get_column_letter(len(columns) + 2)}2")
    wb.save(target)


# ---------------------------------------------------------------- чтение и навигация

@registry.add("read_file", "Read a file (text, code, CSV, Excel, Word/RTF).", {"path": ("string", "")}, ["path"])
def read_file(ctx, path: str) -> str:
    target = resolve(ctx, path)
    if not target.exists():
        return f"Файл не найден: {target}"
    ext = target.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(target, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets[:3]:
            out.append(f"# Лист «{ws.title}»")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 200:
                    out.append("…")
                    break
                out.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    if ext in (".docx", ".doc", ".rtf", ".rtfd", ".odt", ".webarchive"):
        return run(["textutil", "-convert", "txt", "-stdout", str(target)], timeout=30)[:8000]
    data = target.read_bytes()[:200_000]
    return data.decode("utf-8", errors="replace")[:8000]


@registry.add("list_dir", "List a folder (newest first).", {"path": ("string", "default = Coulson's folder")})
def list_dir(ctx, path: str = "") -> str:
    target = resolve(ctx, path) if path else workspace(ctx)
    if not target.is_dir():
        return f"Папка не найдена: {target}"
    items = []
    for p in target.iterdir():
        if p.name.startswith("."):
            continue
        try:
            st = p.stat()
        except OSError:
            continue
        items.append((st.st_mtime, p, st.st_size))
    items.sort(reverse=True)
    lines = [f"{target}:"]
    for mtime, p, size in items[:60]:
        kind = "📁" if p.is_dir() else _human(size / 1024)
        lines.append(f"{p.name} — {kind}, {time.strftime('%d.%m.%Y %H:%M', time.localtime(mtime))}")
    return "\n".join(lines) if items else f"{target}: пусто"


@registry.add("disk_usage", "Disk total/used/free and biggest folders in a path (GB).",
              {"path": ("string", "default home"), "top": ("integer", "")})
def disk_usage(ctx, path: str = "~", top: int = 15) -> str:
    target = resolve(ctx, path or "~")
    st = os.statvfs(str(target))
    total, free = st.f_blocks * st.f_frsize, st.f_bavail * st.f_frsize
    lines = [f"Диск: всего {total / 1e9:.0f} ГБ, занято {(total - free) / 1e9:.0f} ГБ, свободно {free / 1e9:.0f} ГБ"]
    try:
        out = subprocess.run(["du", "-k", "-d", "1", "-x", str(target)], capture_output=True, text=True,
                             timeout=90).stdout
        partial = ""
    except subprocess.TimeoutExpired as e:
        out = e.stdout.decode() if isinstance(e.stdout, bytes) else (e.stdout or "")
        partial = " (подсчёт прерван по времени — данные неполные)"
    sizes = []
    for line in out.splitlines():
        kb, _, p = line.partition("\t")
        if p and Path(p).resolve() != target and kb.isdigit():
            sizes.append((int(kb), Path(p).name))
    sizes.sort(reverse=True)
    lines.append(f"Крупнейшие папки в {target}{partial}:")
    lines += [f"{name}\t{kb / 1024 / 1024:.2f} ГБ" for kb, name in sizes[:max(1, int(top))]]
    return "\n".join(lines)


@registry.add("find_files", "Find files by name (Spotlight).", {"query": ("string", ""), "folder": ("string", "")},
              ["query"])
def find_files(ctx, query: str, folder: str = "") -> str:
    base = str(resolve(ctx, folder)) if folder else str(HOME)
    out = run(["mdfind", "-onlyin", base, "-name", query], timeout=20)
    lines = [l for l in out.splitlines() if "/Library/" not in l and "/." not in l][:20]
    return "\n".join(lines) or "Ничего не найдено"


@registry.add("open_path", "Open a file/folder (reveal = show in Finder).",
              {"path": ("string", ""), "reveal": ("boolean", "")}, ["path"])
def open_path(ctx, path: str, reveal: bool = False) -> str:
    target = resolve(ctx, path)
    return run(["open", "-R", str(target)]) if reveal else _open(ctx, target)
